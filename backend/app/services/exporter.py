from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import datetime

from bidi.algorithm import get_display
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from app.models.entities import Course
from app.services import storage
from app.services.paths import export_path
from app.services.scoring import _get_submission_for_student, compute_gradebook_rows
from app.utils.text import sanitize_filename


@dataclass
class ExportResult:
    csv_path: str
    zip_location: str  # storage location (local path or supabase://...) of the PDFs zip
    report_count: int


def _hebrew(text: str) -> str:
    return get_display(text)


def export_gradebook_csv(db: Session, course_id: int) -> str:
    course = db.get(Course, course_id)
    if not course:
        raise ValueError("Course not found")

    rows = compute_gradebook_rows(db, course_id)
    out_path = export_path(f"{course.code}_{course.year}_{course.semester}", "gradebook.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if rows:
        headers = list(rows[0].keys())
    else:
        headers = [
            "student_id",
            "national_id",
            "name",
            "attendance_percent",
            "attendance_score",
            "hw1",
            "hw2",
            "project_code",
            "project_report",
            "presentation",
            "final_talk",
            "project_base",
            "bonus_points",
            "project_with_bonus",
            "final",
            "overrides",
        ]

    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)

    return str(out_path)


def export_grades_report_xlsx(db: Session, course_id: int) -> str:
    """One row per student: national id, name, final grade, student-facing feedback
    (taken from the project submission). Saved as XLSX."""
    course = db.get(Course, course_id)
    if not course:
        raise ValueError("Course not found")

    rows = compute_gradebook_rows(db, course_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    headers = ["ID", "Name", "Final", "Feedback"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        project = _get_submission_for_student(db, course_id, row["student_id"], "PROJECT")
        feedback = (project.feedback_student or "") if project else ""
        ws.append([row["national_id"], row["name"], row["final"], feedback])

    # Reasonable widths and let feedback wrap.
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 80
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"

    out_path = export_path(f"{course.code}_{course.year}_{course.semester}", "grades_report.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return str(out_path)


def _name_parts(name: str) -> tuple[str, str]:
    parts = [p for p in name.strip().split() if p]
    if not parts:
        return "Unknown", "Student"
    if len(parts) == 1:
        return parts[0], ""
    return parts[-1], " ".join(parts[:-1])


def export_student_reports_hebrew(db: Session, course_id: int) -> ExportResult:
    """Render one Hebrew PDF per student and store them as a single zip through
    `services.storage`, so the hosted site can serve it (signed URL on Supabase)."""
    course = db.get(Course, course_id)
    if not course:
        raise ValueError("Course not found")

    rows = compute_gradebook_rows(db, course_id)
    course_dir_name = f"{course.code}_{course.year}_{course.semester}"

    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for row in rows:
            national_id = row["national_id"]
            last, first = _name_parts(row["name"])
            filename = sanitize_filename(f"{national_id}_{last}_{first}") + ".pdf"

            pdf = FPDF()
            pdf.add_page()
            pdf.add_font("DejaVu", "", font_path)
            pdf.set_font("DejaVu", size=14)

            lines = [
                f"דו\"ח ציונים לסטודנט/ית",
                f"קורס: {course.name} ({course.code}) {course.year} {course.semester}",
                f"שם: {row['name']}",
                f"תעודת זהות: {national_id}",
                f"נוכחות: {row['attendance_score']}",
                f"תרגיל 1: {row['hw1']}",
                f"תרגיל 2: {row['hw2']}",
                f"קוד פרויקט: {row['project_code']}",
                f"דו\"ח פרויקט: {row['project_report']}",
                f"מצגת: {row['presentation']}",
                f"מצגת סופית: {row['final_talk']}",
                f"בונוס: {row['bonus_points']}",
                f"ציון סופי: {row['final']}",
            ]

            for line in lines:
                pdf.cell(0, 10, _hebrew(line), new_x="LMARGIN", new_y="NEXT", align="R")

            archive.writestr(filename, bytes(pdf.output()))

    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%S%f")
    zip_location = storage.save(
        f"exports/{course_dir_name}/student_reports_{stamp}.zip", buffer.getvalue(), "application/zip"
    )

    csv_path = export_gradebook_csv(db, course_id)
    return ExportResult(csv_path=csv_path, zip_location=zip_location, report_count=len(rows))
