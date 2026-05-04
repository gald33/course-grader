from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass

from openpyxl import load_workbook

from app.utils.text import coerce_score_to_0_5, normalize_id


@dataclass
class PresentationRow:
    national_id: str
    stage: str
    clarity: float | None = None
    complexity: float | None = None
    qa_answers: float | None = None
    elegance: float | None = None
    subjective_impression: float | None = None
    bonus: float | None = None


CRITERIA_MAP = {
    "בהירות": "clarity",
    "מורכבות": "complexity",
    "תשובה על שאלות": "qa_answers",
    "תשובות על שאלות": "qa_answers",
    "אלגנט": "elegance",
    "אלגנטיות": "elegance",
    "הערכה כללית": "subjective_impression",
    "subjective impression": "subjective_impression",
    "בונוס": "bonus",
}
IGNORED_CRITERIA = {"נתונים"}


def _normalize_criterion(value: str) -> str:
    return " ".join(str(value).strip().lower().split())


def _maybe_student_id(value: object) -> str:
    sid = normalize_id(value)
    return sid if len(sid) >= 7 else ""


def _extract_sheet_rows(path: str, sheet_name: str, stage: str) -> list[PresentationRow]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    parsed: dict[str, PresentationRow] = {}
    row_idx = 1
    max_col = ws.max_column

    while row_idx <= ws.max_row:
        values = [ws.cell(row=row_idx, column=c).value for c in range(1, max_col + 1)]
        id_columns = [(i, _maybe_student_id(v)) for i, v in enumerate(values) if _maybe_student_id(v)]

        if len(id_columns) >= 1:
            ids = {idx: sid for idx, sid in id_columns}
            for sid in ids.values():
                parsed.setdefault(sid, PresentationRow(national_id=sid, stage=stage))

            cursor = row_idx + 1
            while cursor <= ws.max_row:
                criterion_values = [ws.cell(row=cursor, column=c).value for c in range(1, max_col + 1)]
                next_id_columns = [_maybe_student_id(v) for v in criterion_values if _maybe_student_id(v)]
                if len(next_id_columns) >= 2:
                    break

                label = ""
                for candidate in criterion_values[:5]:
                    if candidate is not None and str(candidate).strip() != "":
                        label = str(candidate).strip()
                        break
                if not label:
                    cursor += 1
                    continue

                if any(x in label for x in IGNORED_CRITERIA):
                    cursor += 1
                    continue

                normalized = _normalize_criterion(label)
                target_key = None
                for raw_key, mapped in CRITERIA_MAP.items():
                    if _normalize_criterion(raw_key) in normalized:
                        target_key = mapped
                        break
                if target_key is None:
                    cursor += 1
                    continue

                for col_idx, sid in ids.items():
                    cell_value = criterion_values[col_idx] if col_idx < len(criterion_values) else None
                    if target_key == "bonus":
                        try:
                            parsed[sid].bonus = float(cell_value) if cell_value not in (None, "") else None
                        except (TypeError, ValueError):
                            parsed[sid].bonus = None
                    else:
                        score = coerce_score_to_0_5(cell_value)
                        if score is not None:
                            setattr(parsed[sid], target_key, score)

                cursor += 1

            row_idx = cursor
            continue

        row_idx += 1

    return list(parsed.values())


def parse_presentations_xlsx(path: str) -> list[PresentationRow]:
    wb = load_workbook(path, data_only=True)
    if len(wb.sheetnames) < 2:
        raise ValueError("Presentation workbook must contain at least 2 sheets (initial and final)")

    initial_sheet = wb.sheetnames[0]
    final_sheet = wb.sheetnames[1]

    rows = _extract_sheet_rows(path, initial_sheet, "initial")
    rows.extend(_extract_sheet_rows(path, final_sheet, "final"))
    return rows


def aggregate_presentation_rows(rows: list[PresentationRow]) -> list[PresentationRow]:
    grouped: dict[tuple[str, str], list[PresentationRow]] = defaultdict(list)
    for row in rows:
        grouped[(row.national_id, row.stage)].append(row)

    aggregated: list[PresentationRow] = []
    fields = ["clarity", "complexity", "qa_answers", "elegance", "subjective_impression", "bonus"]
    for (national_id, stage), items in grouped.items():
        result = PresentationRow(national_id=national_id, stage=stage)
        for field in fields:
            values = [getattr(item, field) for item in items if getattr(item, field) is not None]
            if not values:
                continue
            if field == "bonus":
                setattr(result, field, float(sum(values)))
            else:
                setattr(result, field, float(sum(values) / len(values)))
        aggregated.append(result)

    return aggregated
