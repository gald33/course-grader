from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from app.utils.text import normalize_id


@dataclass
class RosterStudentRow:
    national_id: str
    name: str
    attendance_percent: float | None
    email: str | None = None


ID_HEADER = "תעודת זהות"
NAME_HEADER = "שם"
# Yedion exports split the name and put a title row above the headers.
FIRST_NAME_HEADER = "שם פרטי"
LAST_NAME_HEADER = "שם משפחה"


def _find_header_row(rows: list[tuple]) -> int:
    for i, row in enumerate(rows[:10]):
        if any(cell is not None and str(cell).strip() == ID_HEADER for cell in row):
            return i
    raise ValueError(f"Missing required column '{ID_HEADER}'")


def _email_column(headers: list[str]) -> int | None:
    """Prefer the external (personal) address when both are present."""
    candidates = [i for i, h in enumerate(headers) if "דואר" in h or "מייל" in h or "email" in h.lower()]
    external = [i for i in candidates if "חיצוני" in headers[i]]
    return (external or candidates or [None])[0]


def _clean_email(value: Any) -> str | None:
    text = str(value).strip().lower() if value is not None else ""
    return text if "@" in text and " " not in text else None


def _is_present_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return float(value) > 0
    text = str(value).strip().lower()
    if text in {"", "0", "false", "no", "absent", "לא"}:
        return False
    return text in {"1", "v", "x", "yes", "present", "כן", "✓"}


def parse_roster_xlsx(path: str) -> list[RosterStudentRow]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = _find_header_row(rows)
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_row]]
    id_idx = headers.index(ID_HEADER)
    name_idx = headers.index(NAME_HEADER) if NAME_HEADER in headers else None
    first_idx = headers.index(FIRST_NAME_HEADER) if FIRST_NAME_HEADER in headers else None
    last_idx = headers.index(LAST_NAME_HEADER) if LAST_NAME_HEADER in headers else None
    if name_idx is None and first_idx is None:
        raise ValueError(f"Missing required column '{NAME_HEADER}'")
    email_idx = _email_column(headers)

    attendance_pct_idx = None
    for i, h in enumerate(headers):
        h_low = h.lower()
        if "%" in h or "אחוז" in h or "percent" in h_low:
            attendance_pct_idx = i
            break

    attendance_candidate_indexes = [
        i
        for i, h in enumerate(headers)
        if i not in {id_idx, name_idx} and ("attendance" in h.lower() or "נוכחות" in h or "מפגש" in h)
    ]

    def cell(row: tuple, idx: int | None) -> str:
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    parsed: list[RosterStudentRow] = []
    for row in rows[header_row + 1 :]:
        national_id = normalize_id(row[id_idx] if id_idx < len(row) else "")
        if name_idx is not None:
            name = cell(row, name_idx)
        else:
            name = " ".join(part for part in (cell(row, first_idx), cell(row, last_idx)) if part)
        if not national_id:
            continue

        attendance_percent = None
        if attendance_pct_idx is not None and attendance_pct_idx < len(row):
            raw = row[attendance_pct_idx]
            if raw is not None and str(raw).strip() != "":
                value = float(raw)
                if value <= 1:
                    value *= 100
                attendance_percent = max(0.0, min(100.0, value))

        if attendance_percent is None and attendance_candidate_indexes:
            present = 0
            total = 0
            for idx in attendance_candidate_indexes:
                if idx >= len(row):
                    continue
                cell = row[idx]
                if cell is None or str(cell).strip() == "":
                    continue
                total += 1
                if _is_present_value(cell):
                    present += 1
            if total > 0:
                attendance_percent = round((present / total) * 100, 2)

        parsed.append(
            RosterStudentRow(
                national_id=national_id,
                name=name or national_id,
                attendance_percent=attendance_percent,
                email=_clean_email(row[email_idx]) if email_idx is not None and email_idx < len(row) else None,
            )
        )

    return parsed
